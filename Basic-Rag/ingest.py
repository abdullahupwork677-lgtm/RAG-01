import os
import pdfplumber
import chromadb
from sentence_transformers import SentenceTransformer

# ─── Config ───────────────────────────────────────────────
DOCS_FOLDER = "docs"            # Put your PDFs/txt files here
CHUNK_SIZE   = 500              # Characters per chunk
CHUNK_OVERLAP = 50              # Overlap between chunks (avoids cutting context)
COLLECTION_NAME = "my_docs"
ALLOWED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".pptx"}
# ──────────────────────────────────────────────────────────

_embedder = None


def _get_model():
    global _embedder
    if _embedder is None:
        _embedder = SentenceTransformer("all-MiniLM-L6-v2")
    return _embedder


def _get_collection(client):
    try:
        return client.get_collection(COLLECTION_NAME)
    except Exception:
        return client.create_collection(COLLECTION_NAME)


# ─── Loaders ──────────────────────────────────────────────

def load_pdf_pages(filepath):
    """Extract PDF text page-by-page. Returns [(page_number, text), ...]."""
    pages = []
    with pdfplumber.open(filepath) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                pages.append((i, page_text.strip()))
    return pages


def load_docx(filepath):
    """Read a .docx document (paragraphs + tables) as plain text."""
    from docx import Document
    doc = Document(filepath)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def load_pptx_pages(filepath):
    """Read a .pptx deck. Returns [(slide_number, text), ...]."""
    from pptx import Presentation
    prs = Presentation(filepath)
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        texts.append(line.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        texts.append(" | ".join(cells))
        if texts:
            pages.append((i, "\n".join(texts)))
    return pages


# ─── Chunking ─────────────────────────────────────────────

def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """Split text into overlapping chunks."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks


def _prepare_chunks(filepath, extra_meta=None, user_id=None):
    """Return (pages, chunks, ids, metadatas) for a single file."""
    filename = os.path.basename(filepath)
    ext = os.path.splitext(filename)[1].lower()

    page_units = []
    if ext == ".pdf":
        page_units = load_pdf_pages(filepath)
    elif ext == ".pptx":
        page_units = load_pptx_pages(filepath)
    elif ext in (".txt", ".md"):
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        if text.strip():
            page_units = [(0, text.strip())]
    elif ext == ".docx":
        text = load_docx(filepath)
        if text.strip():
            page_units = [(0, text.strip())]

    prefix = f"{user_id}::" if user_id else ""
    pages, chunks, ids, metas = [], [], [], []
    for page, text in page_units:
        for i, chunk in enumerate(chunk_text(text)):
            pages.append(page)
            chunks.append(chunk)
            ids.append(f"{prefix}{filename}::p{page}::c{i}")
            meta = {"source": filename, "page": int(page), "chunk_index": i}
            if user_id:
                meta["user_id"] = user_id
            if extra_meta:
                meta.update(extra_meta)
            metas.append(meta)
    return pages, chunks, ids, metas


# ─── Ingest ───────────────────────────────────────────────

def ingest_file(filepath, extra_meta=None, user_id=None):
    """Index a single document into ChromaDB (replacing any prior copy of it)."""
    filename = os.path.basename(filepath)
    pages, chunks, ids, metas = _prepare_chunks(filepath, extra_meta, user_id=user_id)
    if not chunks:
        return {"filename": filename, "chunks": 0, "pages": 0, "message": "No readable text found."}

    print(f"  Encoding {len(chunks)} chunks for '{filename}'...")
    embeddings = _get_model().encode(chunks, show_progress_bar=False).tolist()

    client = chromadb.PersistentClient(path="./chroma_db")
    collection = _get_collection(client)
    if user_id:
        collection.delete(where={"$and": [{"source": {"$eq": filename}}, {"user_id": {"$eq": user_id}}]})
    else:
        collection.delete(where={"source": filename})
    collection.add(documents=chunks, embeddings=embeddings, ids=ids, metadatas=metas)

    return {
        "filename": filename,
        "chunks": len(chunks),
        "pages": len(set(pages)),
        "message": f"{len(chunks)} chunks stored for '{filename}'",
    }


def delete_source(filename, user_id=None):
    """Remove every chunk belonging to a source file (optionally for one user)."""
    client = chromadb.PersistentClient(path="./chroma_db")
    collection = _get_collection(client)
    source = os.path.basename(filename)
    if user_id:
        collection.delete(where={"$and": [{"source": {"$eq": source}}, {"user_id": {"$eq": user_id}}]})
    else:
        collection.delete(where={"source": source})


def ingest():
    """Legacy CLI entry: index every supported file inside DOCS_FOLDER."""
    print(f"[1/3] Loading documents from '{DOCS_FOLDER}/'...")
    files = [
        os.path.join(DOCS_FOLDER, name)
        for name in os.listdir(DOCS_FOLDER)
        if os.path.splitext(name)[1].lower() in ALLOWED_EXTENSIONS
    ]

    if not files:
        msg = f"No supported files found in '{DOCS_FOLDER}/'. Supported: "
        msg += ", ".join(sorted(ALLOWED_EXTENSIONS))
        print(f"  {msg}")
        return {"success": False, "message": msg, "chunks": 0, "files": 0}

    total_chunks = 0
    for path in files:
        result = ingest_file(path)
        total_chunks += result["chunks"]
        print(f"  -> {result['message']}")

    msg = f"{total_chunks} chunks stored across {len(files)} files."
    print(f"\nDone! {msg}")
    return {"success": True, "message": msg, "chunks": total_chunks, "files": len(files)}


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Index documents into ChromaDB")
    parser.add_argument("--file", help="Index a single file instead of the whole docs/ folder")
    args = parser.parse_args()

    if args.file:
        print(f"Ingesting single file: {args.file}")
        if not os.path.exists(args.file):
            print(f"  File not found: {args.file}")
        else:
            print(" ", ingest_file(args.file)["message"])
    else:
        ingest()